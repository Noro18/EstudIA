"""
Document processor: file -> list of cleaned Pages (with real page numbers for PDFs).

Responsibilities (nothing more):
  1. Open the file and pull out text, page by page.
  2. Clean it: join hard-wrapped lines, drop page numbers and repeated headers/footers.
  3. Detect scanned (image-only) PDFs and fail with a clear error code.

Chunking lives in chunker.py. process_document() at the bottom wires both together.
"""
from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from .chunker import chunk_pages
from .models import Page, ProcessedDocument

SUPPORTED_TYPES = {".pdf", ".docx", ".txt"}
MIN_PAGE_CHARS = 20          # a page with less text than this counts as "empty"
SCANNED_RATIO = 0.5          # if >= 50% of pages are empty, treat the file as scanned
PSEUDO_PAGE_CHARS = 3000     # DOCX/TXT have no pages, so we make "sections" of this size


class ExtractionError(Exception):
    """Raised with a machine-readable code (matches the PRD's error codes)."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class Extraction:
    pages: list[Page]
    page_count: int
    empty_pages: list[int] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------- cleaning

# "12", "Page 12", "Pág. 12", "Hal. 12", "12 / 40" ... lines that are only a page number
_PAGE_NUMBER_ONLY = re.compile(
    r"(?:p[áa]g(?:ina)?\.?|hal(?:aman|\.)?|page)?\s*\d{1,4}(?:\s*(?:/|de|of)\s*\d{1,4})?",
    re.IGNORECASE,
)


def clean_block(raw: str) -> str:
    """Turn one raw text block (with line breaks from the PDF layout) into a paragraph."""
    text = raw.replace("\u00ad", "")            # soft hyphens
    # Keep the hyphen when a line ends with one. Tetum uses real hyphens a lot
    # (loron-loron, sai-sai), so we never glue "loron-" + "loron" into "loronloron".
    text = re.sub(r"-\s*\n\s*", "-", text)
    text = re.sub(r"\s*\n\s*", " ", text)       # line breaks inside a paragraph -> space
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    return text.strip()


def _normalize_for_repeat_check(text: str) -> str:
    return re.sub(r"\d+", "#", text.lower()).strip()


def drop_repeated_edges(raw_pages: list[tuple[int, list[str]]], min_ratio: float = 0.4):
    """Remove running headers/footers: short first/last paragraphs repeated on many pages.
    Digits are ignored, so 'Relatóriu 2025 - 3' and 'Relatóriu 2025 - 4' count as equal."""
    if len(raw_pages) < 4:
        return raw_pages
    counts: Counter[str] = Counter()
    for _, paras in raw_pages:
        edges = {_normalize_for_repeat_check(p) for p in (paras[:1] + paras[-1:]) if len(p) < 120}
        counts.update(edges)
    threshold = max(3, int(len(raw_pages) * min_ratio))
    repeated = {t for t, c in counts.items() if c >= threshold}

    cleaned = []
    for number, paras in raw_pages:
        kept = []
        for i, p in enumerate(paras):
            is_edge = i == 0 or i == len(paras) - 1
            if is_edge and _normalize_for_repeat_check(p) in repeated:
                continue
            kept.append(p)
        cleaned.append((number, kept))
    return cleaned


def _is_noise(paragraph: str) -> bool:
    return len(paragraph) <= 14 and bool(_PAGE_NUMBER_ONLY.fullmatch(paragraph.strip()))


# --------------------------------------------------------------------------- readers

def _read_pdf(path: Path) -> list[tuple[int, list[str]]]:
    try:
        import pymupdf
    except ImportError:  # very old PyMuPDF
        import fitz as pymupdf

    try:
        doc = pymupdf.open(path)
    except Exception as exc:
        raise ExtractionError("extraction_failed", f"Could not open PDF: {exc}") from exc

    if doc.needs_pass:
        raise ExtractionError("extraction_failed", "The PDF is password protected.")

    raw_pages = []
    for number, page in enumerate(doc, start=1):
        # "blocks" gives paragraph-like units in reading order: (x0, y0, x1, y1, text, no, type)
        blocks = [b for b in page.get_text("blocks", sort=True) if b[6] == 0]  # 0 = text block
        paras = [clean_block(b[4]) for b in blocks]
        raw_pages.append((number, [p for p in paras if p]))
    doc.close()
    return raw_pages


def _paginate(paragraphs: list[str], chars_per_page: int = PSEUDO_PAGE_CHARS):
    """DOCX/TXT have no page numbers. Group paragraphs into numbered 'sections' instead."""
    pages, current, size = [], [], 0
    for p in paragraphs:
        if current and size + len(p) > chars_per_page:
            pages.append((len(pages) + 1, current))
            current, size = [], 0
        current.append(p)
        size += len(p)
    if current:
        pages.append((len(pages) + 1, current))
    return pages


def _read_txt(path: Path) -> list[tuple[int, list[str]]]:
    text = path.read_text(encoding="utf-8", errors="replace")
    paragraphs = [clean_block(p) for p in re.split(r"\n\s*\n", text)]
    return _paginate([p for p in paragraphs if p])


def _read_docx(path: Path) -> list[tuple[int, list[str]]]:
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise ExtractionError("extraction_failed", "python-docx is not installed.") from exc
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ExtractionError("extraction_failed", f"Could not open DOCX: {exc}") from exc
    paragraphs = [clean_block(p.text) for p in document.paragraphs]
    return _paginate([p for p in paragraphs if p])


# --------------------------------------------------------------------------- public API

def extract_document(path: str | Path) -> Extraction:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_TYPES:
        raise ExtractionError("unsupported_type", f"Unsupported file type: {suffix or '(none)'}")

    warnings: list[str] = []
    if suffix == ".pdf":
        raw_pages = _read_pdf(path)
    elif suffix == ".docx":
        raw_pages = _read_docx(path)
        warnings.append("DOCX has no real page numbers; 'pages' are ~3000-character sections.")
    else:
        raw_pages = _read_txt(path)
        warnings.append("TXT has no real page numbers; 'pages' are ~3000-character sections.")

    raw_pages = drop_repeated_edges(raw_pages)
    pages = [
        Page(number, [p for p in paras if not _is_noise(p)])
        for number, paras in raw_pages
    ]

    page_count = len(pages)
    if page_count == 0:
        raise ExtractionError("empty_document", "The document has no pages.")

    empty = [p.number for p in pages if p.char_count < MIN_PAGE_CHARS]
    if len(empty) / page_count >= SCANNED_RATIO:
        raise ExtractionError(
            "scanned_no_text",
            "This file seems to be scanned images with no selectable text. OCR is required.",
        )
    if empty:
        warnings.append(f"Pages with (almost) no text, possibly scanned: {empty}")

    return Extraction(pages=pages, page_count=page_count, empty_pages=empty, warnings=warnings)


def process_document(
    path: str | Path,
    doc_id: str | None = None,
    max_chars: int = 1200,
    overlap_chars: int = 150,
    min_chars: int = 300,
) -> ProcessedDocument:
    """Extract + chunk in one call. This is what Django's ingestion job will call."""
    path = Path(path)
    doc_id = doc_id or re.sub(r"[^a-z0-9]+", "-", path.stem.lower()).strip("-") or "doc"

    extraction = extract_document(path)
    chunks = chunk_pages(
        extraction.pages,
        doc_id=doc_id,
        max_chars=max_chars,
        overlap_chars=overlap_chars,
        min_chars=min_chars,
    )
    return ProcessedDocument(
        doc_id=doc_id,
        filename=path.name,
        page_count=extraction.page_count,
        empty_pages=extraction.empty_pages,
        warnings=extraction.warnings,
        chunks=chunks,
    )
